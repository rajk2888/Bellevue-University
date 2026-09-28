"""Render the recorded deck (recorded.json) onto the SEMICON West 2026 template.

Usage: python3 apply-template.py TEMPLATE.pptx recorded.json OUT.pptx

recorded.json comes from running build-deck.js with a recording stand-in for
pptxgenjs. This script keeps the template's master, layouts, logos, and fonts,
maps the recorded content into the template's content area, and swaps the
palette to the SEMI brand colors.
"""
import base64
import copy
import io
import json
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt
from lxml import etree

TEMPLATE, RECORDED, OUT = sys.argv[1:4]

# Template content area (template is 26.667 x 15 in)
L, R = 2.57, 24.0
K = (R - L) / (12.733 - 0.6)  # scale from the recorded 13.333 x 7.5 in canvas
CT = 3.6                      # where recorded y=1.65 lands
FS = 1.5                      # font scale (Gotham runs wider than Calibri)

HEAD, BODY, BLACK = "Gotham Medium", "Gotham Book", "Gotham Black"
OTEXT = "B35900"              # orange dark enough for text on white

COLOR = {
    "1B2229": "1A1A1A", "27303A": "2E2E2E", "C46A2B": "E97B05", "F6E6DA": "FDEBD3",
    "5A6672": "5E5E5E", "8A96A3": "8C8C8C", "EEF1F4": "F2F2F2", "D5DBE1": "D5D5D5",
    "2F7F7A": "3F7F12", "DDEFEC": "E6F2DD", "A63D32": "A31F23", "F6E1DE": "F8E0E0",
    "E8B48F": "F4AE00", "3A4552": "3A3A3A", "FFFFFF": "FFFFFF", "2B7FB8": "2B7FB8",
}


def col(c):
    return COLOR.get((c or "000000").upper(), c)


def X(x):
    return L + (x - 0.6) * K


def Y(y):
    return CT + (y - 1.65) * K


def D(d):
    return d * K


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation(TEMPLATE)
layouts = {l.name: l for l in prs.slide_layouts}
tpl_slides = list(prs.slides)
title_slide, close_slide = tpl_slides[0], tpl_slides[4]
slide_num_sp = next(sh._element for sh in tpl_slides[2].placeholders
                    if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)
slide_num_sp = copy.deepcopy(slide_num_sp)

# Drop the template's sample slides 2-4
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    part = prs.part.related_part(sldId.rId)
    if part in (tpl_slides[1].part, tpl_slides[2].part, tpl_slides[3].part):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)


def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def set_run(run, o, size, big=False):
    f = run.font
    head = o.get("fontFace") == "Calibri" and o.get("bold")
    f.name = BLACK if big else (HEAD if head else BODY)
    f.size = Pt(size)
    f.bold = False
    f.italic = bool(o.get("italic"))
    c = col(o.get("color", "1B2229"))
    if (o.get("color") or "").upper() == "C46A2B" and not big:
        c = OTEXT
    f.color.rgb = rgb(c)
    if o.get("charSpacing"):
        run._r.get_or_add_rPr().set("spc", str(int(o["charSpacing"] * 100)))
    if o.get("hyperlink"):
        run.hyperlink.address = o["hyperlink"]["url"]


def add_text(slide, t, o, box=None):
    x, y, w, h = box or (X(o["x"]), Y(o["y"]), D(o["w"]), D(o["h"]))
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(o.get("valign"), MSO_ANCHOR.TOP)
    size = o.get("fontSize", 14) * FS
    big = o.get("fontSize", 14) >= 30
    runs = t if isinstance(t, list) else [{"text": t, "options": {}}]
    p = tf.paragraphs[0]
    for r in runs:
        ro = {**o, **(r.get("options") or {})}
        run = p.add_run()
        run.text = r["text"]
        set_run(run, ro, size, big)
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(o.get("align"), PP_ALIGN.LEFT)
        if o.get("paraSpaceAfter"):
            p.space_after = Pt(o["paraSpaceAfter"] * FS)
        if ro.get("breakLine"):
            p = tf.add_paragraph()
    return tb


SHAPES = {"RECTANGLE": MSO_SHAPE.RECTANGLE, "ROUNDED_RECTANGLE": MSO_SHAPE.ROUNDED_RECTANGLE, "CHEVRON": MSO_SHAPE.CHEVRON}


def add_shape(slide, type_, o):
    x, y, w, h = X(o["x"]), Y(o["y"]), D(o["w"]), D(o["h"])
    line = o.get("line") or {}
    if type_ == "LINE":
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y + h))
        strip_style(c)
        c.line.color.rgb = rgb(col(line.get("color")))
        c.line.width = Pt(line.get("width", 1) * FS)
        return
    sh = slide.shapes.add_shape(SHAPES[type_], Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(sh)
    fill = (o.get("fill") or {}).get("color")
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(col(fill))
    else:
        sh.fill.background()
    if line.get("color") and line.get("color") != fill:
        sh.line.color.rgb = rgb(col(line["color"]))
        sh.line.width = Pt(line.get("width", 1) * FS)
        if line.get("dashType") == "dash":
            sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    if type_ == "ROUNDED_RECTANGLE" and o.get("rectRadius"):
        sh.adjustments[0] = min(0.5, o["rectRadius"] / min(o["w"], o["h"]))


def recolor(data):
    from PIL import Image
    im = Image.open(io.BytesIO(data)).convert("RGBA")
    px = im.load()
    changed = False
    for yy in range(im.height):
        for xx in range(im.width):
            r_, g_, b_, a_ = px[xx, yy]
            if a_ and (r_, g_, b_) == (0x2F, 0x7F, 0x7A):
                px[xx, yy] = (0x54, 0xA4, 0x1C, a_)
                changed = True
    if not changed:
        return data
    out = io.BytesIO()
    im.save(out, "PNG")
    return out.getvalue()


def add_image(slide, o):
    data = recolor(base64.b64decode(o["data"].split("base64,", 1)[1]))
    slide.shapes.add_picture(io.BytesIO(data), Inches(X(o["x"])), Inches(Y(o["y"])), Inches(D(o["w"])), Inches(D(o["h"])))


def add_chart(slide, data, o):
    cd = CategoryChartData()
    cd.categories = data[0]["labels"]
    for s in data:
        cd.add_series(s["name"], s["values"])
    ctype = XL_CHART_TYPE.COLUMN_CLUSTERED if o.get("barDir") == "col" else XL_CHART_TYPE.BAR_CLUSTERED
    gf = slide.shapes.add_chart(ctype, Inches(X(o["x"])), Inches(Y(o["y"])), Inches(D(o["w"])), Inches(D(o["h"])), cd)
    ch = gf.chart
    ch.font.name = BODY
    plot = ch.plots[0]
    plot.gap_width = o.get("barGapWidthPct", 100)
    colors = [col(c) for c in o.get("chartColors", [])]
    if o.get("varyColors"):
        plot.vary_by_categories = True
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
    else:
        plot.vary_by_categories = False
        for i, ser in enumerate(plot.series):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
    if o.get("showValue"):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = o.get("dataLabelFormatCode", "General")
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(o.get("dataLabelFontSize", 12) * FS)
        dl.font.name = HEAD if o.get("dataLabelFontBold") else BODY
        dl.font.color.rgb = rgb(col(o.get("dataLabelColor", "1B2229")))
    va = ch.value_axis
    va.visible = not o.get("valAxisHidden")
    va.has_major_gridlines = False
    if "valAxisMaxVal" in o:
        va.maximum_scale = o["valAxisMaxVal"]
    if "valAxisMinVal" in o:
        va.minimum_scale = o["valAxisMinVal"]
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(o.get("catAxisLabelFontSize", 12) * FS)
    ca.tick_labels.font.color.rgb = rgb(col(o.get("catAxisLabelColor", "1B2229")))
    ca.reverse_order = o.get("catAxisOrientation") == "maxMin"
    if o.get("catAxisLineShow") is False:
        ca.format.line.fill.background()
    else:
        ca.format.line.color.rgb = rgb(col(o.get("catAxisLineColor", "D5DBE1")))
    ch.has_legend = bool(o.get("showLegend"))
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(o.get("legendFontSize", 12) * FS)
        ch.legend.font.color.rgb = rgb(col(o.get("legendColor", "1B2229")))
    if o.get("showTitle"):
        ch.has_title = True
        tf = ch.chart_title.text_frame
        tf.text = o["title"]
        f = tf.paragraphs[0].runs[0].font
        f.size = Pt(o.get("titleFontSize", 14) * FS)
        f.name = HEAD
        f.bold = False
        f.color.rgb = rgb(col(o.get("titleColor", "5A6672")))
    else:
        ch.has_title = False


def cell_border(cell, color="D5D5D5", w=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w=str(w))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)


def add_table(slide, rows, o):
    nr, nc = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nr, nc, Inches(X(o["x"])), Inches(Y(o["y"])), Inches(D(o["w"])), Inches(D(sum(o["rowH"]))))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, w in enumerate(o["colW"]):
        tbl.columns[j].width = Inches(D(w))
    for i, h in enumerate(o["rowH"]):
        tbl.rows[i].height = Inches(D(h))
    for i, row in enumerate(rows):
        for j, c in enumerate(row):
            co = c["options"]
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(col(co["fill"]["color"]))
            cell.margin_left = cell.margin_right = Inches(D(0.15))
            cell.margin_top = cell.margin_bottom = Inches(D(0.06))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            run.text = c["text"]
            set_run(run, {**co, "fontFace": "Calibri" if co.get("bold") else "x"}, co.get("fontSize", 13) * FS)
            cell_border(cell)


def content_slide(rec):
    s = prs.slides.add_slide(layouts["Title & Text"])
    body = [ph for ph in s.placeholders if ph.placeholder_format.idx == 1]
    for ph in body:
        ph._element.getparent().remove(ph._element)
    s.shapes._spTree.append(copy.deepcopy(slide_num_sp))
    num = s.shapes[-1]
    num.left, num.top, num.width, num.height = Inches(24.3), Inches(13.33), Inches(1.14), Inches(0.39)
    for it in rec["items"]:
        o = it.get("o") or {}
        if it["k"] == "text":
            t = it["t"]
            txt = t if isinstance(t, str) else ""
            if txt.startswith("SEMICON West  |"):
                continue
            if o.get("y") == 0.2 and o.get("fontSize") == 11:          # kicker
                add_text(s, t, {**o, "fontSize": 13.5}, box=(L, 1.45, R - L, 0.45))
                continue
            if o.get("y") == 0.45 and o.get("fontSize") in (34,):      # title
                title = s.shapes.title
                title.left, title.top, title.width, title.height = Inches(L), Inches(1.95), Inches(R - L), Inches(1.3)
                title.text_frame.text = t
                title.text_frame.word_wrap = True
                title.text_frame.paragraphs[0].runs[0].font.size = Pt(46)
                continue
            add_text(s, t, o)
        elif it["k"] == "shape":
            add_shape(s, it["type"], o)
        elif it["k"] == "image":
            add_image(s, o)
        elif it["k"] == "chart":
            add_chart(s, it["data"], o)
        elif it["k"] == "table":
            add_table(s, it["rows"], o)
    if rec.get("notes"):
        s.notes_slide.notes_text_frame.text = rec["notes"]
    return s


def section_slide(title, sub):
    s = prs.slides.add_slide(layouts["Section"])
    s.shapes.title.text = title
    add_text(s, sub, {"fontSize": 18, "color": "5A6672"}, box=(1.18, 4.95, 12.2, 2.4))
    s.notes_slide.notes_text_frame.text = "Section divider. Click straight through."
    return s


def white_text(slide, text, box, size, font=HEAD):
    tb = slide.shapes.add_textbox(*[Inches(v) for v in box])
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        p.alignment = PP_ALIGN.LEFT
        r.font.name, r.font.size, r.font.color.rgb = font, Pt(size), rgb("FFFFFF")


rec = json.load(open(RECORDED))

# Title slide: reuse the template's own slide 1
for sh in title_slide.shapes:
    if sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
        sh.left, sh.top, sh.width, sh.height = Inches(0.64), Inches(2.55), Inches(14.2), Inches(4.3)
        tf = sh.text_frame
        tf.text = "Addressing the risks associated with the global semiconductor supply chain"
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        tf.paragraphs[0].runs[0].font.size = Pt(64)
    elif sh.name.startswith("Subtitle"):
        sh.top = Inches(7.15)
        sh.text_frame.paragraphs[0].runs[0].text = "Leveraging Data and AI Governance"
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r.text = ""
    elif sh.has_text_frame:
        sh.top, sh.width = Inches(8.1), Inches(14.2)
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "[Your Name], [Title, Company]  |  [Session date]"
        p.runs[0].font.size = Pt(32)
        for r in p.runs[1:]:
            r.text = ""
title_slide.notes_slide.notes_text_frame.text = rec[0]["notes"]

# Close slide: reuse the template's "Thank you"
white_text(close_slide, "See your risk sooner. Act on it with data you can defend.", (0.64, 8.45, 13.8, 1.2), 34)
white_text(close_slide, "Questions?", (0.64, 9.75, 13.8, 1.0), 44)
white_text(close_slide, "[Your Name]  |  [email or LinkedIn URL]", (0.64, 10.95, 13.8, 0.8), 28, BODY)
close_slide.notes_slide.notes_text_frame.text = rec[15]["notes"]

order = [title_slide, content_slide(rec[1])]
order.append(section_slide("The risk map", "Where the global chain is fragile, and why it matters now"))
order += [content_slide(rec[i]) for i in range(2, 7)]
order.append(section_slide("Why data is the gap", "The data exists. Trusting and connecting it is the problem."))
order += [content_slide(rec[i]) for i in range(7, 9)]
order.append(section_slide("Governance as the lever", "Controls that turn risk data into decisions you can defend"))
order += [content_slide(rec[i]) for i in range(9, 15)]
order.append(close_slide)
order += [content_slide(rec[16]), content_slide(rec[17])]

by_part = {prs.part.related_part(e.rId): e for e in sldIdLst}
els = [by_part[s.part] for s in order]
for e in list(sldIdLst):
    sldIdLst.remove(e)
for e in els:
    sldIdLst.append(e)

# Rename slide and notes parts in final order so no two parts share a name
for i, sl in enumerate(order, 1):
    sl.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
    if sl.has_notes_slide:
        sl.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{i}.xml")

prs.save(OUT)
print("wrote", OUT, len(els), "slides")
